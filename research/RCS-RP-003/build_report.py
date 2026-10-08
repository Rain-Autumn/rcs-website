from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "generated" / "RCS-RP-003_Report_FR_v2.0.docx"
GREEN = "315F45"
MUTED = "64746A"


def rows_from_csv(name: str) -> list[dict[str, str]]:
    with (ROOT / "generated" / name).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def repeat_table_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    props.append(repeat)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def add_markdown_runs(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|https?://\S+)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1]); run.font.name = "Consolas"; run.font.size = Pt(9)
        elif part.startswith("http"):
            clean = part.rstrip(".,)")
            run = paragraph.add_run(clean); run.font.color.rgb = RGBColor(49, 95, 69); run.underline = True
        else:
            paragraph.add_run(part)


def add_data_table(doc: Document, title: str, selected: list[dict[str, str]]) -> None:
    doc.add_paragraph(title, style="Caption")
    headers = ["Jeu / échantillon", "Méthode", "Entrée logique (octets)", "Référence (octets)", "Sortie médiane (octets)", "Variation médiane", "Essais"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    table.autofit = True
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade(cell, GREEN)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True; run.font.color.rgb = RGBColor(255, 255, 255); run.font.size = Pt(8)
    repeat_table_header(table.rows[0])
    for item in selected:
        cells = table.add_row().cells
        name = item["Sample"]
        if item["Dataset"] != "Synthetic":
            name = f"{item['Dataset']} · {name}"
        values = [
            name,
            item["Method"],
            f"{int(item['InputLogicalBytes']):,}".replace(",", " "),
            f"{float(item['ReferenceAllocatedBytesMean']):,.0f}".replace(",", " "),
            f"{float(item['OutputBytesMedian']):,.0f}".replace(",", " "),
            f"{float(item['ChangePercentMedian']):+.2f}%",
            item["RepeatCount"],
        ]
        for i, value in enumerate(values):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[i].paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                for run in paragraph.runs:
                    run.font.size = Pt(7.5)


def build() -> None:
    source = (ROOT / "study-manuscript.md").read_text(encoding="utf-8")
    cases = rows_from_csv("summary-by-case.csv")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7); section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72); section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.32); section.footer_distance = Inches(0.3)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 27), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11)):
        styles[name].font.name = "Aptos Display"; styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(GREEN)
        styles[name].font.bold = True
    styles["Caption"].font.name = "Aptos"; styles["Caption"].font.size = Pt(8); styles["Caption"].font.color.rgb = RGBColor.from_string(MUTED)

    header = section.header.paragraphs[0]
    header.text = "RAIJU CLOUD SYSTEM   /   RESEARCH PROGRAM   /   RCS-RP-003"
    header.style = "Caption"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("RCS-RP-003 · Version 2.0     |     ")
    add_page_number(footer)

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(80)
    run = p.add_run("RAIJU CLOUD SYSTEM  /  RESEARCH"); run.bold = True; run.font.size = Pt(11); run.font.color.rgb = RGBColor.from_string(GREEN)
    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(24)
    title.add_run("Libérer de l’espace disque\nsans perdre ses données")
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(9)
    add_markdown_runs(p, "Rapport de recherche · RCS-RP-003 · Version 2.0")
    for line in ("8 octobre 2026", "Contributeur : Hugues Henrotte", "DOI réservé : 10.5281/zenodo.23233639 — dépôt non publié"):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        r = p.add_run(line); r.font.color.rgb = RGBColor.from_string(MUTED)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(75)
    p.add_run("Étude indépendante du benchmark RCS-RP-002. Terminal personnel Windows · Volume NTFS · Aucune donnée personnelle analysée.").italic = True
    doc.add_page_break()

    in_body = False
    inserted_aggregates = inserted_synthetic = inserted_appendix = False
    for raw in source.splitlines():
        line = raw.strip()
        if line == "## Résumé":
            in_body = True
        if not in_body or not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("### Résumé global par corpus"):
            doc.add_heading("Résumé global par corpus", level=3)
            selected = [r for r in cases if r["Sample"] == "Complete corpus folder"]
            add_data_table(doc, "Tableau 1 — Essais agrégés par corpus; référence propre à chaque méthode.", selected)
            inserted_aggregates = True
            continue
        if line.startswith("### Cas synthétiques"):
            doc.add_heading("Cas synthétiques", level=3)
            selected = [r for r in cases if r["Dataset"] == "Synthetic"]
            add_data_table(doc, "Tableau 2 — Cinq motifs générés de façon déterministe.", selected)
            inserted_synthetic = True
            continue
        if line.startswith("## Annexe A"):
            doc.add_page_break()
            doc.add_heading("Annexe A — Résultats détaillés par fichier", level=1)
            public = [r for r in cases if r["Dataset"] != "Synthetic" and r["Sample"] != "Complete corpus folder"]
            add_data_table(doc, "Tableau A1 — Médianes par entrée publique et méthode (3 répétitions chacune).", public)
            doc.add_paragraph("La colonne « variation » compare la sortie à la référence de la même méthode. Pour les valeurs brutes, les plages, écarts-types et empreintes, consulter le classeur de données et les CSV joints.")
            inserted_appendix = True
            continue
        if line.startswith("Les tableaux exhaustifs sont produits"):
            continue
        if line.startswith(("Le générateur de rapport insère ici", "Le générateur insère ici", "Les résultats par fichier sont générés")):
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3); continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1); continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_markdown_runs(p, line[2:]); continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_markdown_runs(p, re.sub(r"^\d+\. ", "", line)); continue
        p = doc.add_paragraph()
        add_markdown_runs(p, line)

    if not (inserted_aggregates and inserted_synthetic and inserted_appendix):
        raise RuntimeError(f"Could not inject generated result tables: aggregates={inserted_aggregates}, synthetic={inserted_synthetic}, appendix={inserted_appendix}.")
    fig = ROOT / "generated" / "figures" / "RCS-RP-003-Silesia-ratio.png"
    if fig.exists():
        doc.add_page_break(); doc.add_heading("Figure A1 — Corpus Silesia", level=1)
        doc.add_paragraph("Médiane de trois essais par fichier; un ratio supérieur à 1 indique une sortie plus petite.")
        doc.add_picture(str(fig), width=Inches(6.7))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "RCS-RP-003 — Libérer de l’espace disque sans perdre ses données"
    doc.core_properties.author = "Raiju Cloud System"
    doc.core_properties.subject = "Étude de stockage Windows, compression NTFS, ZIP et hygiène de données"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
