from __future__ import annotations

import argparse
import csv
import math
import statistics
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    # Windows PowerShell exports floating-point values with the current regional decimal separator.
    for row in rows:
        for key in ("ChangePercent", "RatioInputOverOutput", "DurationMilliseconds"):
            if key in row and "," in row[key] and "." not in row[key]:
                row[key] = row[key].replace(",", ".")
    return rows


def write_csv(path: Path, rows: list[dict], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def median(values: list[float]) -> float:
    return statistics.median(values)


def sample_sd(values: list[float]) -> float:
    return statistics.stdev(values) if len(values) > 1 else 0.0


def make_summaries(trials: list[dict[str, str]]) -> list[dict]:
    groups: dict[tuple[str, str, str, str], list[dict[str, str]]] = {}
    for row in trials:
        groups.setdefault((row["Dataset"], row["Sample"], row["Category"], row["Method"]), []).append(row)

    case_rows: list[dict] = []
    for (dataset, sample, category, method), rows in sorted(groups.items()):
        output = [int(row["OutputBytes"]) for row in rows]
        change = [float(row["ChangePercent"]) for row in rows]
        duration = [float(row["DurationMilliseconds"]) for row in rows]
        logical = [int(row["InputLogicalBytes"]) for row in rows]
        baseline = [int(row["ReferenceAllocatedBytes"]) for row in rows]
        hashes_ok = all(row["IntegrityPassed"].lower() == "true" for row in rows)
        case_rows.append(
            {
                "Dataset": dataset,
                "Sample": sample,
                "Category": category,
                "Method": method,
                "RepeatCount": len(rows),
                "FileCount": int(rows[0]["Files"]),
                "InputLogicalBytes": logical[0],
                "ReferenceAllocatedBytesMean": round(statistics.mean(baseline), 3),
                "OutputBytesMean": round(statistics.mean(output), 3),
                "OutputBytesMedian": round(median(output), 3),
                "OutputBytesMin": min(output),
                "OutputBytesMax": max(output),
                "OutputBytesSampleSD": round(sample_sd(output), 3),
                "ChangePercentMean": round(statistics.mean(change), 6),
                "ChangePercentMedian": round(median(change), 6),
                "ChangePercentMin": round(min(change), 6),
                "ChangePercentMax": round(max(change), 6),
                "ChangePercentSampleSD": round(sample_sd(change), 6),
                "RatioInputOverOutputMedian": round(median([float(row["RatioInputOverOutput"]) for row in rows]), 6),
                "DurationMillisecondsMedian": round(median(duration), 3),
                "IntegrityPassed": hashes_ok,
            }
        )

    return case_rows


def get_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def draw_silesia_chart(case_rows: list[dict], target: Path) -> None:
    silesia = [row for row in case_rows if row["Dataset"] == "Silesia" and row["Sample"] != "Complete corpus folder"]
    by_case: dict[str, dict[str, float]] = defaultdict(dict)
    for row in silesia:
        by_case[row["Sample"]][row["Method"]] = float(row["RatioInputOverOutputMedian"])
    names = list(by_case)
    width, height = 1900, 1000
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font, label_font, small_font = get_font(34, True), get_font(20), get_font(16)
    draw.text((95, 35), "Silesia corpus: input-to-output ratio by file", fill="#202b25", font=title_font)
    draw.text((95, 78), "Median of three runs; higher ratio means a smaller archive / allocated footprint", fill="#52665b", font=label_font)

    left, right, top, bottom = 105, 1840, 160, 830
    max_value = max(max(values.values()) for values in by_case.values()) if by_case else 1
    axis_max = max(2, math.ceil(max_value / 2) * 2)
    draw.line((left, top, left, bottom), fill="#6c7a70", width=2)
    draw.line((left, bottom, right, bottom), fill="#6c7a70", width=2)
    for tick in range(0, axis_max + 1, max(1, axis_max // 6)):
        y = bottom - (bottom - top) * tick / axis_max
        draw.line((left, int(y), right, int(y)), fill="#e0e5e0", width=1)
        draw.text((55, int(y) - 12), str(tick), fill="#52665b", font=small_font)

    colors = {"ZIP Optimal": "#3b7057", "NTFS compact": "#d6a64a"}
    group_w = (right - left) / max(1, len(names))
    bar_w = min(40, group_w * 0.28)
    for index, name in enumerate(names):
        center = left + group_w * (index + 0.5)
        for offset, method in ((-bar_w * 0.6, "ZIP Optimal"), (bar_w * 0.6, "NTFS compact")):
            value = by_case[name].get(method, 0)
            bar_h = value / axis_max * (bottom - top)
            x0 = center + offset - bar_w / 2
            x1 = x0 + bar_w
            y0 = bottom - bar_h
            draw.rounded_rectangle((int(x0), int(y0), int(x1), bottom), radius=4, fill=colors[method])
            draw.text((int(x0) - 2, int(y0) - 25), f"{value:.2f}", fill="#273b32", font=small_font)
        label = name if len(name) <= 12 else name[:11] + "…"
        bbox = draw.textbbox((0, 0), label, font=small_font)
        draw.text((int(center - (bbox[2] - bbox[0]) / 2), bottom + 14), label, fill="#273b32", font=small_font)

    legend_y = 900
    for index, method in enumerate(("ZIP Optimal", "NTFS compact")):
        x = 680 + index * 330
        draw.rectangle((x, legend_y, x + 24, legend_y + 22), fill=colors[method])
        draw.text((x + 34, legend_y - 2), method, fill="#273b32", font=label_font)
    target.parent.mkdir(parents=True, exist_ok=True)
    image.save(target, optimize=True)


def draw_change_chart(case_rows: list[dict], target: Path) -> None:
    selected = [row for row in case_rows if row["Dataset"] == "Silesia" and row["Sample"] != "Complete corpus folder"]
    names = list(dict.fromkeys(row["Sample"] for row in selected))
    values = {name: {} for name in names}
    for row in selected:
        values[row["Sample"]][row["Method"]] = float(row["ChangePercentMedian"])
    width, height = 1900, 1100
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font, label_font, small_font = get_font(32, True), get_font(18), get_font(15)
    draw.text((85, 30), "Silesia corpus: median change in measured bytes", fill="#202b25", font=title_font)
    draw.text((85, 72), "ZIP compares archive bytes with logical input; NTFS compares allocated bytes before/after", fill="#52665b", font=label_font)
    left, right, top, bottom = 450, 1830, 145, 965
    all_values = [value for group in values.values() for value in group.values()]
    min_v, max_v = min(all_values, default=-100), max(all_values, default=0)
    low = min(-100, math.floor(min_v / 10) * 10)
    high = max(20, math.ceil(max_v / 10) * 10)
    span = max(1, high - low)
    zero_x = left + (0 - low) / span * (right - left)
    draw.line((int(zero_x), top, int(zero_x), bottom), fill="#58685e", width=2)
    row_h = (bottom - top) / max(1, len(names))
    for idx, name in enumerate(names):
        cy = top + row_h * (idx + 0.5)
        label = name if len(name) < 35 else name[:32] + "…"
        draw.text((25, int(cy - 10)), label, fill="#273b32", font=small_font)
        for offset, method, color in ((-7, "ZIP Optimal", "#3b7057"), (8, "NTFS compact", "#d6a64a")):
            val = values[name].get(method, 0)
            x = left + (val - low) / span * (right - left)
            draw.line((int(zero_x), int(cy + offset), int(x), int(cy + offset)), fill=color, width=8)
            draw.ellipse((int(x - 5), int(cy + offset - 5), int(x + 5), int(cy + offset + 5)), fill=color)
    for tick in range(int(low // 20 * 20), int(high) + 1, 20):
        if tick < low or tick > high:
            continue
        x = left + (tick - low) / span * (right - left)
        draw.line((int(x), bottom, int(x), bottom + 8), fill="#6c7a70", width=1)
        draw.text((int(x - 20), bottom + 15), f"{tick}%", fill="#52665b", font=small_font)
    for index, (method, color) in enumerate((("ZIP Optimal", "#3b7057"), ("NTFS compact", "#d6a64a"))):
        x = 760 + index * 320
        draw.rectangle((x, 1035, x + 23, 1054), fill=color)
        draw.text((x + 32, 1033), method, fill="#273b32", font=label_font)
    target.parent.mkdir(parents=True, exist_ok=True)
    image.save(target, optimize=True)


def add_sheet(workbook: Workbook, title: str, rows: list[dict]) -> None:
    sheet = workbook.create_sheet(title)
    if not rows:
        sheet.append(["No data"])
        return
    headers = list(rows[0].keys())
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315F45")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    integer_fields = {
        "Repeat", "Files", "RepeatCount", "FileCount", "InputLogicalBytes",
        "ReferenceAllocatedBytes", "ReferenceAllocatedBytesMean", "OutputBytes",
        "OutputBytesMean", "OutputBytesMedian", "OutputBytesMin", "OutputBytesMax",
        "ChangeBytes", "ExpectedBytes", "ActualBytes",
    }
    float_fields = {
        "ChangePercent", "RatioInputOverOutput", "DurationMilliseconds",
        "OutputBytesSampleSD", "ChangePercentMean", "ChangePercentMedian",
        "ChangePercentMin", "ChangePercentMax", "ChangePercentSampleSD",
        "RatioInputOverOutputMedian", "DurationMillisecondsMedian",
    }
    for row in rows:
        values = []
        for header in headers:
            value = row.get(header)
            if value not in (None, "") and header in integer_fields:
                value = int(float(value))
            elif value not in (None, "") and header in float_fields:
                value = float(value)
            elif header == "IntegrityPassed" and isinstance(value, str):
                value = value.lower() == "true"
            values.append(value)
        sheet.append(values)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for col_idx, header in enumerate(headers, 1):
        max_len = len(str(header))
        for cell in list(sheet.columns)[col_idx - 1][1:201]:
            max_len = max(max_len, min(64, len(str(cell.value or ""))))
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 42)
    for row_idx in range(2, sheet.max_row + 1):
        if row_idx % 2 == 0:
            for cell in sheet[row_idx]:
                cell.fill = PatternFill("solid", fgColor="F4F7F4")
        for cell in sheet[row_idx]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            header = headers[cell.column - 1]
            if header.endswith("Bytes") or header.endswith("BytesMean") or header.endswith("BytesMedian") or header.endswith("BytesMin") or header.endswith("BytesMax") or header in {"Repeat", "Files", "RepeatCount", "FileCount", "ExpectedBytes", "ActualBytes"}:
                cell.number_format = "#,##0"
            elif "Percent" in header:
                cell.number_format = "0.00\"%\""
            elif "Ratio" in header:
                cell.number_format = "0.0000"
            elif "Milliseconds" in header:
                cell.number_format = "#,##0.0"


def build_workbook(output: Path, trials: list[dict], case_rows: list[dict], sources: list[dict] | None, checks: list[dict]) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Overview"
    first.append(["RCS-RP-003 — Résultats reproductibles"])
    first.merge_cells("A1:H1")
    first["A1"].font = Font(name="Aptos Display", size=16, bold=True, color="315F45")
    for row in [
        ["Population", "26 fichiers publics · 225 908 846 octets · Canterbury, Canterbury Large, Silesia"],
        ["Mesures", "220 lignes d’essais · 110 exécutions appariées · ZIP Optimal et compression NTFS"],
        ["Répétitions", "3 par fichier public · 5 par cas synthétique · agrégats: 3 / 3 / 1"],
        ["Intégrité", "Tous les contrôles de source et de transformation réussis (246 au total)"],
        ["Interprétation", "ZIP compare à la taille logique; NTFS compare à l’espace alloué avant compression."],
        ["Limites", "Un système NTFS; pas de fichiers personnels, de benchmark de vitesse ou de recommandation de suppression."],
        [],
        ["Agrégats des corpus (médianes de sortie)"],
        ["Corpus", "Méthode", "Fichiers", "Octets logiques", "Référence (octets)", "Sortie médiane", "Variation médiane", "Répétitions"],
    ]:
        first.append(row)
    aggregate_rows = [row for row in case_rows if row["Sample"] == "Complete corpus folder"]
    for row in aggregate_rows:
        first.append([
            row["Dataset"], row["Method"], row["FileCount"], row["InputLogicalBytes"],
            row["ReferenceAllocatedBytesMean"], row["OutputBytesMedian"],
            row["ChangePercentMedian"], row["RepeatCount"],
        ])
    first.append([])
    first.append(["Cas synthétiques (médianes de sortie)"])
    first.append(["Cas", "Méthode", "Octets logiques", "Référence (octets)", "Sortie médiane", "Variation médiane", "Répétitions"])
    for row in case_rows:
        if row["Dataset"] == "Synthetic":
            first.append([
                row["Sample"], row["Method"], row["InputLogicalBytes"],
                row["ReferenceAllocatedBytesMean"], row["OutputBytesMedian"],
                row["ChangePercentMedian"], row["RepeatCount"],
            ])
    for row_idx in range(2, 8):
        first.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=8)
        first.cell(row_idx, 2).alignment = Alignment(vertical="center", wrap_text=True)
        first.row_dimensions[row_idx].height = 30
    first.column_dimensions["A"].width = 35
    for col in "BCDEFGH":
        first.column_dimensions[col].width = 22
    for row_idx in (9, 10, 18, 19):
        for cell in first[row_idx]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="315F45")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row_idx in list(range(11, 17)) + list(range(20, first.max_row + 1)):
        for cell in first[row_idx]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, (int, float)):
                is_percent = (11 <= row_idx <= 16 and cell.column == 7) or (row_idx >= 20 and cell.column == 6)
                cell.number_format = "0.00\"%\"" if is_percent else "#,##0"
    for row_idx in range(1, first.max_row + 1):
        if first.cell(row_idx, 1).value in {"Agrégats des corpus (médianes de sortie)", "Cas synthétiques (médianes de sortie)"}:
            first.cell(row_idx, 1).font = Font(bold=True, color="315F45", size=12)
    first.freeze_panes = "A10"
    add_sheet(workbook, "Case summary", case_rows)
    add_sheet(workbook, "Trial results", trials)
    add_sheet(workbook, "Corpus checks", checks)
    if sources:
        add_sheet(workbook, "Sources", sources)
    workbook.save(output)


def write_summary_markdown(path: Path, case_rows: list[dict]) -> None:
    lines = [
        "# Calculated results (generated from results.csv)", "",
        "ZIP changes use logical input bytes; NTFS changes use pre-compression allocated bytes.", "",
        "## Corpus aggregate cases", "",
        "| Dataset | Method | Files | Logical bytes | Reference bytes | Output median | Change median | Ratio | Runs | Integrity |",
        "|---|---|---:|---:|---:|---:|---:|---:|---:|---|",
    ]
    for row in case_rows:
        if row["Sample"] != "Complete corpus folder":
            continue
        lines.append(
            f"| {row['Dataset']} | {row['Method']} | {row['FileCount']} | {row['InputLogicalBytes']} | "
            f"{row['ReferenceAllocatedBytesMean']:.0f} | {row['OutputBytesMedian']:.0f} | "
            f"{row['ChangePercentMedian']:.4f}% | {row['RatioInputOverOutputMedian']:.4f} | {row['RepeatCount']} | {row['IntegrityPassed']} |"
        )
    lines.extend(["", "## Synthetic cases", "", "| Case | Method | Logical bytes | Reference bytes | Output median | Change median | Ratio | Runs | Integrity |", "|---|---|---:|---:|---:|---:|---:|---:|---|"])
    for row in case_rows:
        if row["Dataset"] != "Synthetic":
            continue
        lines.append(
            f"| {row['Sample']} | {row['Method']} | {row['InputLogicalBytes']} | {row['ReferenceAllocatedBytesMean']:.0f} | "
            f"{row['OutputBytesMedian']:.0f} | {row['ChangePercentMedian']:.4f}% | {row['RatioInputOverOutputMedian']:.4f} | {row['RepeatCount']} | {row['IntegrityPassed']} |"
        )
    lines.extend(["", "## Public corpus files", "", "| Dataset | File | Method | Logical bytes | Reference bytes | Output median | Change median | Ratio | Runs | Integrity |", "|---|---|---|---:|---:|---:|---:|---:|---:|---|"])
    for row in case_rows:
        if row["Dataset"] == "Synthetic" or row["Sample"] == "Complete corpus folder":
            continue
        lines.append(
            f"| {row['Dataset']} | {row['Sample']} | {row['Method']} | {row['InputLogicalBytes']} | {row['ReferenceAllocatedBytesMean']:.0f} | "
            f"{row['OutputBytesMedian']:.0f} | {row['ChangePercentMedian']:.4f}% | {row['RatioInputOverOutputMedian']:.4f} | {row['RepeatCount']} | {row['IntegrityPassed']} |"
        )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=ROOT / "generated")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "generated")
    args = parser.parse_args()
    trial_path = args.input_dir / "results.csv"
    if not trial_path.exists():
        raise FileNotFoundError(trial_path)
    trials = read_csv(trial_path)
    if not trials:
        raise ValueError("The trial results file is empty.")
    if any(row["IntegrityPassed"].lower() != "true" for row in trials):
        raise ValueError("At least one content-integrity check failed; analysis aborted.")

    case_rows = make_summaries(trials)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_dir / "summary-by-case.csv", case_rows, list(case_rows[0]))
    figures = args.output_dir / "figures"
    draw_silesia_chart(case_rows, figures / "RCS-RP-003-Silesia-ratio.png")
    draw_change_chart(case_rows, figures / "RCS-RP-003-Silesia-change.png")

    source_path = args.input_dir / "sources.csv"
    if not source_path.exists():
        source_path = ROOT / "sources.csv"
    checks_path = args.input_dir / "source-checks.csv"
    sources = read_csv(source_path) if source_path.exists() else None
    checks = read_csv(checks_path) if checks_path.exists() else []
    build_workbook(args.output_dir / "RCS-RP-003-data.xlsx", trials, case_rows, sources, checks)
    write_summary_markdown(args.output_dir / "results-summary.md", case_rows)

    print(f"Trial rows: {len(trials)}")
    print(f"Case/method summaries: {len(case_rows)}")
    print(f"Datasets represented: {len({row['Dataset'] for row in trials})}")
    print(f"Integrity checks: {sum(row['IntegrityPassed'].lower() == 'true' for row in trials)}/{len(trials)}")


if __name__ == "__main__":
    main()
